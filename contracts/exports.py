"""Выгрузки реестра договоров в .xlsx и .docx."""
from io import BytesIO

from django.http import HttpResponse
from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Font


def contracts_to_xlsx(contracts):
    wb = Workbook()
    ws = wb.active
    ws.title = "Договоры"
    headers = ["Номер", "Обучающийся", "Программа", "Статус", "Сумма", "Оплачено", "Начало"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for contract in contracts:
        ws.append([
            contract.number,
            str(contract.student),
            contract.program.name,
            contract.status.name,
            contract.amount_rub,
            contract.paid_rub,
            contract.starts_on.strftime("%d.%m.%Y"),
        ])
    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="contracts.xlsx"'
    return response


def contracts_to_docx(contracts):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    doc.add_heading("Реестр договоров на платные образовательные услуги", level=1)
    doc.add_paragraph("Веб-сервис «Договор.Онлайн». Автор: Курбатов Дмитрий Евгеньевич.")
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for index, title in enumerate(["Номер", "Обучающийся", "Программа", "Статус", "Сумма, руб."]):
        table.rows[0].cells[index].text = title
    for contract in contracts:
        row = table.add_row().cells
        row[0].text = contract.number
        row[1].text = str(contract.student)
        row[2].text = contract.program.name
        row[3].text = contract.status.name
        row[4].text = str(contract.amount_rub)
    buffer = BytesIO()
    doc.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    response["Content-Disposition"] = 'attachment; filename="contracts.docx"'
    return response
